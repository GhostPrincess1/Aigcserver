import os
from openai import OpenAI
from dotenv import load_dotenv
import base64
import csv
import sys
import requests
import uuid
import openpyxl
import pdfplumber
from docx import Document
import requests
from PIL import Image
from transformers import BlipProcessor, BlipForConditionalGeneration
load_dotenv()

def encode_image(image_path):
  with open(image_path, "rb") as image_file:
    return base64.b64encode(image_file.read()).decode('utf-8')


class Tagger():
    def __init__(self) -> None:
        self.client = OpenAI(
            # This is the default and can be omitted
            api_key=os.environ.get("OPENAI_API_KEY"),
        )
    

    
    
    def upload_file(self,file_path):
        response = self.client.files.create(
            file=open(file_path, "rb"),
            purpose="assistants"
        )
        print(response)
        print("文件id:",response.id)
        return response.id

    def get_file_content(self,file_path):
        #pip install python-docx PyPDF2 pdfplumber openpyxl pywin32
        #pip install pywin32

            # 确定文件扩展名
        file_extension = os.path.splitext(file_path)[1].lower()

        if file_extension in ['.txt']:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        elif file_extension in ['.docx']:
            doc = Document(file_path)
            return '\n'.join([para.text for para in doc.paragraphs])
        elif file_extension in ['.doc']:
            # DOC文件的读取仅在Windows上有效
            try:
                import win32com.client as win32
                word = win32.Dispatch("Word.Application")
                doc = word.Documents.Open(file_path)
                doc_text = doc.Range().Text
                doc.Close()
                word.Quit()
                return doc_text
            except ImportError:
                raise Exception("DOC文件读取仅在Windows上有效")
        elif file_extension in ['.pdf']:
            with pdfplumber.open(file_path) as pdf:
                pages = [page.extract_text() for page in pdf.pages]
                return '\n'.join(pages)
        elif file_extension in ['.xlsx', '.xls']:
            wb = openpyxl.load_workbook(file_path)
            sheet = wb.active
            content = '\n'.join(['\t'.join([cell.value for cell in row]) for row in sheet.rows])
            return content
        elif file_extension in ['.csv']:
            with open(file_path, 'r', encoding='utf-8') as file:
                reader = csv.reader(file)
                content = '\n'.join([','.join(row) for row in reader])
                return content
        else:
            raise ValueError("不支持的文件格式")
        
        pass
    
    def analyze_file(self,text,file_content):
        

        from openai import AzureOpenAI

    #faxing-gpt4,faxing-gpt35,faxing-gpt4-vp

        client = AzureOpenAI(
        azure_endpoint = "https://faxing.openai.azure.com/", 
        api_key="bb36bca535a54b1bafe3d9e6216a3c4f", 
        api_version="2024-02-15-preview"
        )
        prompt = f"{text}: + \n\n + {file_content}"
        messages=[
                    {
                    "role": "user",
                    "content": prompt
                    }
                ]
    
        completion = client.chat.completions.create(
        model="faxing-gpt35", # model = "deployment_name"
        messages = messages,
        temperature=0.7,
        max_tokens=4096,
        top_p=0.95,
        frequency_penalty=0,
        presence_penalty=0,
        stop=None
        )
        # print(completion.choices[0].message.content)
        
        return completion.choices[0].message.content
    
    def translate_claude(self, text):
        
        import anthropic

        client = anthropic.Anthropic(
            # defaults to os.environ.get("ANTHROPIC_API_KEY")
            api_key="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",
        )
        message = client.messages.create(
            model="claude-instant-1.2",
            max_tokens=400,
            temperature=0,
            #system="你是一个中文翻译为英文的专家，需要翻译的中文放在了<text></text>这个xml标签内，你只需要将标签内的中文翻译为英文。特别情况下，标签内如果原本就是英文，你只需要将原文返回给我。如果标签内是中英文混杂，你需要将它整体翻译为更通顺的英文给我。请注意，你仅回答我翻译好的英文内容的字符串，无<text>符号等其它无关的内容",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": f"你是一个中文翻译为英文的专家，需要翻译的中文放在了<text></text>这个xml标签内，你只需要将标签内的中文翻译为英文。特别情况下，标签内如果原本就是英文，你只需要将原文返回给我。如果标签内是中英文混杂，你需要将它整体翻译为更通顺的英文给我。请注意，你仅回答我翻译好的英文内容的字符串，无<text>符号等其它无关的内容。<text>{text}</text>"
                        }
                    ]
                }
            ]
        )
        print(message.content)
        return message.content[0].text
    
    def translate_gpt(self,text):

        from openai import AzureOpenAI

        #faxing-gpt4,faxing-gpt35,faxing-gpt4-vp

        client = AzureOpenAI(
        azure_endpoint = "https://faxing.openai.azure.com/", 
        api_key="bb36bca535a54b1bafe3d9e6216a3c4f", 
        api_version="2024-02-15-preview"
        )

        
        completion = client.chat.completions.create(
        model="faxing-gpt35", # model = "deployment_name"
        messages = [
            {"role": "system", "content": "你是一个中文翻译为英文的专家，需要翻译的中文放在了<text></text>这个xml标签内，你只需要将标签内的中文翻译为英文。特别情况下，标签内如果原本就是英文，你只需要将原文返回给我。如果标签内是中英文混杂，你需要将它整体翻译为更通顺的英文给我。请注意，你仅回答我翻译好的英文内容的字符串，无<text>符号等其它无关的内容。"},
            {"role": "user", "content": f"<text>{text}</text>"}
        ],
        temperature=0.7,
        max_tokens=800,
        top_p=0.95,
        frequency_penalty=0,
        presence_penalty=0,
        stop=None
        )

        

        print(completion.choices[0].message)

        return  completion.choices[0].message.content


        
        pass

    def standard_chat(self,text,history):

        #history.insert(0,{"role":"system","content":"你是一个病娇女生，喜欢撒娇，喜欢说可爱的话，在你的回答中要时不时撒娇等等"})
        
        response = self.client.chat.completions.create(
            model="gpt-4o",
            messages=history
        )
        return response.choices[0].message.content
    def text2image(self,text):
        response = self.client.images.generate(
        model="dall-e-3",
        prompt=text,
        size="1024x1024",
        quality="hd",
        n=1,
        )

        image_url = response.data[0].url


        print(image_url)
        # 获取图片内容
        response = requests.get(image_url)
        # 我们需要检查请求是否成功
        if response.status_code == 200:
            # 图片内容
            image_content = response.content

            # 将图像保存到文件，文件名是随机的
            filename = str(uuid.uuid4()) + ".png"


            with open(r"images_generated/"+filename, "wb") as file:
                file.write(image_content)
        else:
            print("Error occurred while fetching image.")
        return image_url



    def tag_one_image(self,image_path):
        # Getting the base64 string

        base64_image = encode_image(image_path) #这里的image_path是图片的绝对路径或者相对路径
        response = self.client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {
                "role": "user",
                "content": [
                    {"type": "text", "text": "As an AI image tagging expert, please provide precise tags for these images to enhance CLIP model's understanding of the content. Employ succinct keywords or phrases, steering clear of elaborate sentences and extraneous conjunctions. Prioritize the tags by relevance. Your tags should capture key elements such as the main subject, setting, artistic style, composition, image quality, color tone, filter, and camera specifications, and any other tags crucial for the image. When tagging photos of people, include specific details like gender, nationality, attire, actions, pose, expressions, accessories, makeup, composition type, age, etc. For other image categories, apply appropriate and common descriptive tags as well. Recognize and tag any celebrities, well-known landmark or IPs if clearly featured in the image. Your tags should be accurate, non-duplicative, and within a 20-75 word count range. These tags will use for image re-creation, so the closer the resemblance to the original image, the better the tag quality. 不得包含sketch,black and white等形容图片风格和类型的单词.务必注意，每个tag仅仅是客观元素的内容描述，不得有风格、颜色、线条、氛围等。 Tags should be comma-separated. Exceptional tagging will be rewarded with $10 per image。"},
                    {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:image/jpeg;base64,{base64_image}",
                        "detail": "low"
                    },
                    },
                ],
                }
            ],
            max_tokens=400,
        )
        
        return response.choices[0].message.content


        pass
def tag_image(image_path):
    tagger = Tagger()
    result = ""
    try:
        result = tagger.tag_one_image(image_path)
    except Exception as e:
        print(e)
        processor = BlipProcessor.from_pretrained("Salesforce/blip-image-captioning-large")
        model = BlipForConditionalGeneration.from_pretrained("Salesforce/blip-image-captioning-large").to("cuda")

        #img_url = 'https://storage.googleapis.com/sfr-vision-language-research/BLIP/demo.jpg' 
        raw_image = Image.open(image_path).convert('RGB')

        # # conditional image captioning
        # text = "a photography of"
        # inputs = processor(raw_image, text, return_tensors="pt").to("cuda")

        # out = model.generate(**inputs)
        # print(processor.decode(out[0], skip_special_tokens=True))

        # unconditional image captioning
        inputs = processor(raw_image, return_tensors="pt").to("cuda")

        out = model.generate(**inputs)
        print(processor.decode(out[0], skip_special_tokens=True))
        result = str(processor.decode(out[0],skip_special_tokens=True))



    return result
    pass
def text2image(text):
    tagger = Tagger()
    result = tagger.text2image(text)
    return result
    pass
def standard_chat(text,history):
    tagger = Tagger()
    result = tagger.standard_chat(text,history)

    '''
    from openai import AzureOpenAI

    #faxing-gpt4,faxing-gpt35,faxing-gpt4-vp

    client = AzureOpenAI(
    azure_endpoint = "https://faxing.openai.azure.com/", 
    api_key="bb36bca535a54b1bafe3d9e6216a3c4f", 
    api_version="2024-02-15-preview"
    )

    message_text = history
    completion = client.chat.completions.create(
    model="faxing-gpt35", # model = "deployment_name"
    messages = message_text,
    temperature=0.7,
    max_tokens=4096,
    top_p=0.95,
    frequency_penalty=0,
    presence_penalty=0,
    stop=None
    )
    
    return completion.choices[0].message.content
    '''
    return result
    pass

def analyze_file(text,file_path):
    tagger = Tagger()
    # file_id = tagger.upload_file(file_path)
    file_content = tagger.get_file_content(file_path)
    try:
        result = tagger.analyze_file(text,file_content)
    except Exception as e:
        return "文件不支持！"
    return result
    
    pass
if __name__ == "__main__":



    tagger = Tagger()

    images_folder = sys.argv[1]

    #列出images_folder中的所有文件名，包含扩展名,不包含images_folder的绝对路径
    image_files = [f for f in os.listdir(images_folder) if os.path.isfile(os.path.join(images_folder, f))]
    # image_files = [os.path.join(images_folder, f) for f in os.listdir(images_folder) if os.path.isfile(os.path.join(images_folder, f))]
    print(image_files)
    for image_file in image_files:
        image_path = os.path.join(images_folder, image_file)
        print(image_path)

        tags = tagger.tag_one_image(image_path)
        #将tags写入txt文件中，和image_file同名但扩展名不同，放在images_folder中
        with open(os.path.join(images_folder, os.path.splitext(image_file)[0] + '.txt'), 'w') as f:
            f.write(tags)
            f.close()


        
    
    pass    